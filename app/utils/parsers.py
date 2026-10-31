import csv
from pathlib import Path

import docx
import pdfplumber
from pptx import Presentation

from app.core.exceptions import UnsupportedFileTypeError


def extract_text_from_file(file_path: str, mime_type: str) -> str:
    """
    Synchronously extracts raw text from a document based on its MIME type.
    Should be run in a threadpool (e.g. anyio.to_thread) to avoid blocking the event loop.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    text = ""

    try:
        if mime_type == "application/pdf":
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n\n"

        elif mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
            doc = docx.Document(path)
            for para in doc.paragraphs:
                text += para.text + "\n"

        elif mime_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            prs = Presentation(path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
                text += "\n"

        elif mime_type == "text/csv":
            with open(path, "r", encoding="utf-8") as f:
                reader = csv.reader(f)
                for row in reader:
                    text += " | ".join(row) + "\n"

        elif mime_type in ("text/plain", "text/markdown"):
            with open(path, "r", encoding="utf-8") as f:
                text = f.read()

        else:
            raise UnsupportedFileTypeError(f"Unsupported MIME type for extraction: {mime_type}")

    except Exception as e:
        if isinstance(e, UnsupportedFileTypeError):
            raise
        raise RuntimeError(f"Failed to extract text from {path.name}: {e}") from e

    # Clean up excess whitespace while preserving paragraph structure
    import re
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
